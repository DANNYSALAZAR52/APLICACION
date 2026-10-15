"""Exportador a Excel (.xlsx) y CSV con fórmulas neutralizadas (DOCS.MD §8)."""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from fleetplanner.application.ports.outbound.exports import Sheet
from fleetplanner.domain.masters.schema import neutralize_cell

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def sheet_title(name: str, used: set[str]) -> str:
    """Nombre válido de hoja (máx. 31 caracteres, sin [ ] : * ? / \\) y sin repetirse."""
    base = _BAD_SHEET_CHARS.sub(" ", name).strip()[:31] or "Hoja"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


class SpreadsheetExporter:
    def to_xlsx(self, sheets: Sequence[Sheet]) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        used: set[str] = set()
        for sheet in sheets:
            ws = workbook.create_sheet(sheet_title(sheet.name, used))
            row_no = 1
            for note in sheet.notes:
                ws.cell(row=row_no, column=1, value=neutralize_cell(note)).font = Font(italic=True)
                row_no += 1
            if sheet.notes:
                row_no += 1
            header_row = row_no
            for col, title in enumerate(sheet.header, start=1):
                cell = ws.cell(row=row_no, column=col, value=neutralize_cell(title))
                cell.font = Font(bold=True)
            for values in sheet.rows:
                row_no += 1
                for col, value in enumerate(values, start=1):
                    safe = neutralize_cell(value)
                    cell = ws.cell(row=row_no, column=col, value=safe)
                    if isinstance(safe, str):
                        cell.data_type = "s"  # texto literal: nunca se evalúa como fórmula
            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
            for col in range(1, len(sheet.header) + 1):
                longest = max(
                    [len(str(sheet.header[col - 1]))]
                    + [len(str(r[col - 1])) for r in sheet.rows if col - 1 < len(r)]
                )
                ws.column_dimensions[get_column_letter(col)].width = min(max(10, longest + 2), 60)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def to_csv(self, sheet: Sheet) -> bytes:
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\r\n")
        for note in sheet.notes:
            writer.writerow([neutralize_cell(note)])
        if sheet.notes:
            writer.writerow([])
        writer.writerow([neutralize_cell(h) for h in sheet.header])
        for row in sheet.rows:
            writer.writerow([neutralize_cell(v) for v in row])
        return ("﻿" + out.getvalue()).encode("utf-8")
