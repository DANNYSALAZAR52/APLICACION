"""Exportador a Excel/CSV: fórmulas neutralizadas, nombres de hoja válidos y números intactos."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from fleetplanner.adapters.outbound.export import SpreadsheetExporter
from fleetplanner.adapters.outbound.export.spreadsheets import sheet_title
from fleetplanner.application.ports.outbound.exports import Sheet

EVIL = ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1)", "\tTAB", "\rCR"]


def workbook(*sheets: Sheet):  # type: ignore[no-untyped-def]
    return load_workbook(io.BytesIO(SpreadsheetExporter().to_xlsx(sheets)))


class TestFormulaNeutralization:
    @pytest.mark.parametrize("text", EVIL)
    def test_text_that_could_be_a_formula_is_stored_as_literal_text(self, text: str) -> None:
        ws = workbook(Sheet("Datos", ["Texto"], [[text]])).active
        cell = ws["A2"]
        assert cell.data_type == "s"
        assert cell.value.replace("\n", "\r") == "'" + text  # xlsx normaliza \r a \n

    @pytest.mark.parametrize("text", EVIL)
    def test_csv_neutralizes_the_same_way(self, text: str) -> None:
        raw = SpreadsheetExporter().to_csv(Sheet("Datos", ["Texto"], [[text]])).decode("utf-8-sig")
        assert ("'" + text).replace("\r", "") in raw.replace("\r\n", "\n").replace("\r", "")

    def test_headers_and_notes_are_neutralized_too(self) -> None:
        ws = workbook(Sheet("Datos", ["=cmd"], [[1]], notes=("@nota",))).active
        assert ws["A1"].value == "'@nota" and ws["A3"].value == "'=cmd"

    def test_real_numbers_are_not_altered(self) -> None:
        ws = workbook(Sheet("Datos", ["n"], [[-3], [2.5], [0]])).active
        assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == [-3, 2.5, 0]
        assert ws["A2"].data_type == "n"

    def test_ordinary_text_is_untouched(self) -> None:
        ws = workbook(Sheet("Datos", ["t"], [["Calle 80 · Bogotá"]])).active
        assert ws["A2"].value == "Calle 80 · Bogotá"

    def test_none_cells_stay_empty(self) -> None:
        ws = workbook(Sheet("Datos", ["a", "b"], [[None, 1]])).active
        assert ws["A2"].value is None and ws["B2"].value == 1


class TestStructure:
    def test_accents_survive_in_xlsx_and_csv(self) -> None:
        sheet = Sheet("Proyección", ["Año", "Déficit"], [["Camión", 3]])
        assert workbook(sheet).active["A2"].value == "Camión"
        assert "Camión" in SpreadsheetExporter().to_csv(sheet).decode("utf-8-sig")

    def test_csv_starts_with_a_bom_so_excel_reads_utf8(self) -> None:
        assert SpreadsheetExporter().to_csv(Sheet("x", ["a"], [])).startswith(b"\xef\xbb\xbf")

    def test_header_is_bold_and_panes_are_frozen_below_it(self) -> None:
        ws = workbook(Sheet("Datos", ["a"], [[1]], notes=("nota",))).active
        assert ws["A3"].font.bold and ws.freeze_panes == "A4"

    def test_several_sheets_are_created_in_order(self) -> None:
        wb = workbook(Sheet("Uno", ["a"], []), Sheet("Dos", ["b"], []))
        assert wb.sheetnames == ["Uno", "Dos"]

    def test_sheet_names_are_valid_and_unique(self) -> None:
        used: set[str] = set()
        assert sheet_title("Ruta [A]/B:C*?", used) == "Ruta  A  B C"
        assert sheet_title("ruta [a]/b:c*?", used).endswith("(2)")
        assert len(sheet_title("x" * 80, used)) <= 31
        assert sheet_title("", used) == "Hoja"

    def test_same_title_twice_does_not_break_the_workbook(self) -> None:
        wb = workbook(Sheet("Datos", ["a"], []), Sheet("Datos", ["b"], []))
        assert len(set(wb.sheetnames)) == 2
