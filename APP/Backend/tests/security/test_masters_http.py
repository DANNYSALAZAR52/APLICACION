"""Seguridad y contrato HTTP de maestros: RBAC, aislamiento, CSV, exportaciones y auditoría."""

from __future__ import annotations

import io
import warnings
from dataclasses import dataclass

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fleetplanner.bootstrap.application import Application, SeedCredential, build_http_app
from tenancy_helpers import A, B, make_app, password_of

warnings.filterwarnings("ignore", message=".*httpx.*")

M = "/api/v1/masters"
TYPE = {
    "name": "furgon", "capacity_t": 1.5, "capacity_m3": 9, "pallets": 3,
    "fixed_cost_month": 2_000_000, "variable_cost_per_km": 700, "third_party_daily_rate": 300_000,
}  # fmt: skip
CUSTOMER_HEADER = "code,name,address,lat,lon,service_minutes,window_start,window_end,restrictions"


@dataclass
class Ctx:
    app: Application
    creds: list[SeedCredential]
    client: TestClient

    def headers(self, tenant: str | None, username: str) -> dict[str, str]:
        key = tenant if tenant else "(sin empresa)"
        response = self.client.post(
            "/api/v1/auth/login",
            json={
                "tenant_slug": tenant,
                "username": username,
                "password": password_of(self.creds, key, username),
            },
        )
        assert response.status_code == 200, response.text
        return {"Authorization": f"Bearer {response.json()['access_token']}"}

    def admin(self, tenant: str = A) -> dict[str, str]:
        return self.headers(tenant, "admin")


@pytest.fixture
def ctx() -> Ctx:
    app, creds, _ = make_app()
    return Ctx(app, creds, TestClient(build_http_app(app), raise_server_exceptions=False))


def _msg(response) -> str:  # type: ignore[no-untyped-def]
    return response.json()["error"]["mensaje"]  # type: ignore[no-any-return]


class TestAccess:
    @pytest.mark.parametrize(
        ("method", "path"),
        [
            ("get", "/schema"),
            ("get", "/entities/vehicle_type"),
            ("post", "/entities/vehicle_type"),
            ("put", "/entities/vehicle_type/turbo"),
            ("delete", "/entities/vehicle_type/turbo"),
            ("post", "/entities/customer/csv/preview"),
            ("post", "/entities/customer/csv/apply"),
            ("get", "/entities/customer/export"),
            ("get", "/parameters"),
            ("put", "/parameters/reserve"),
            ("get", "/history"),
        ],
    )
    def test_every_endpoint_requires_a_token(self, ctx: Ctx, method: str, path: str) -> None:
        assert getattr(ctx.client, method)(M + path).status_code == 401

    @pytest.mark.parametrize("username", ["gerencia", "despachador", "conductor"])
    def test_only_the_company_admin_can_manage_masters(self, ctx: Ctx, username: str) -> None:
        h = ctx.headers(A, username)
        assert ctx.client.get(f"{M}/schema", headers=h).status_code == 403
        assert ctx.client.get(f"{M}/entities/customer", headers=h).status_code == 403
        assert (
            ctx.client.post(
                f"{M}/entities/vehicle_type", json={"values": TYPE}, headers=h
            ).status_code
            == 403
        )
        assert (
            ctx.client.put(f"{M}/parameters/reserve", json={"value": 9}, headers=h).status_code
            == 403
        )
        assert ctx.client.get(f"{M}/history", headers=h).status_code == 403
        assert ctx.client.get(f"{M}/entities/customer/export", headers=h).status_code == 403

    def test_the_superadmin_has_no_company_and_cannot_edit_company_data(self, ctx: Ctx) -> None:
        h = ctx.headers(None, "superadmin")
        assert ctx.client.get(f"{M}/entities/customer", headers=h).status_code == 403

    def test_denied_attempts_are_audited(self, ctx: Ctx) -> None:
        ctx.client.post(
            f"{M}/entities/vehicle_type", json={"values": TYPE}, headers=ctx.headers(A, "gerencia")
        )
        assert any(
            e.action == "master.create" and e.outcome.value == "denied"
            for e in ctx.app.audit.list_events(None, 50)
        )

    def test_unknown_entities_and_malformed_paths_are_rejected(self, ctx: Ctx) -> None:
        h = ctx.admin()
        assert ctx.client.get(f"{M}/entities/inventado", headers=h).status_code == 404
        assert ctx.client.get(f"{M}/entities/Cus-tomer", headers=h).status_code == 422
        assert ctx.client.delete(f"{M}/entities/customer/%00", headers=h).status_code in (404, 422)


class TestSchemaAndCrud:
    def test_schema_describes_entities_parameters_and_lookups(self, ctx: Ctx) -> None:
        body = ctx.client.get(f"{M}/schema", headers=ctx.admin()).json()
        assert body["demo_label"] == "DATOS DE DEMOSTRACIÓN"
        names = {e["name"] for e in body["entities"]}
        assert {"vehicle_type", "corridor", "customer", "vehicle", "order", "daily_demand"} <= names
        assert "reserve" in {p["key"] for p in body["parameters"]}
        assert "turbo" in body["lookups"]["vehicle_types"]
        customer = next(e for e in body["entities"] if e["name"] == "customer")
        assert customer["key_fields"] == ["code"]
        assert any(f["name"] == "lat" and f["min"] == -4.3 for f in customer["fields"])

    def test_create_update_delete_roundtrip_with_audit_and_history(self, ctx: Ctx) -> None:
        h = ctx.admin()
        created = ctx.client.post(f"{M}/entities/vehicle_type", json={"values": TYPE}, headers=h)
        assert created.status_code == 201, created.text
        updated = ctx.client.put(
            f"{M}/entities/vehicle_type/furgon", json={"values": {**TYPE, "pallets": 5}}, headers=h
        )
        assert updated.status_code == 200 and updated.json()["pallets"] == 5
        assert ctx.client.delete(f"{M}/entities/vehicle_type/furgon", headers=h).status_code == 204
        history = ctx.client.get(f"{M}/history?entity=vehicle_type&limit=10", headers=h).json()
        assert [c["action"] for c in history[:3]] == ["delete", "update", "create"]
        assert history[1]["before"]["pallets"] == 3 and history[1]["after"]["pallets"] == 5
        assert history[1]["actor"] == "admin"
        actions = {
            e.action for e in ctx.app.audit.list_events(None, 100) if e.actor_username == "admin"
        }
        assert {"master.create", "master.update", "master.delete"} <= actions

    def test_validation_errors_come_back_in_spanish_with_all_problems(self, ctx: Ctx) -> None:
        bad = {**TYPE, "capacity_t": 0, "name": ""}
        response = ctx.client.post(
            f"{M}/entities/vehicle_type", json={"values": bad}, headers=ctx.admin()
        )
        assert response.status_code == 422
        assert "Nombre es obligatorio" in _msg(response) and "mayor que 0" in _msg(response)

    def test_criterion_11_over_http(self, ctx: Ctx) -> None:
        customer = {
            "code": "C-900", "name": "Nuevo", "lat": 40, "lon": -74.1, "service_minutes": 15,
            "window_start": "14:00", "window_end": "10:00",
        }  # fmt: skip
        response = ctx.client.post(
            f"{M}/entities/customer", json={"values": customer}, headers=ctx.admin()
        )
        assert response.status_code == 422
        assert (
            "Latitud" in _msg(response)
            and "14:00–10:00" in _msg(response)
            and "posterior" in _msg(response)
        )

    def test_referenced_records_cannot_be_deleted(self, ctx: Ctx) -> None:
        response = ctx.client.delete(f"{M}/entities/vehicle_type/turbo", headers=ctx.admin())
        assert response.status_code == 409 and "lo usan" in _msg(response)

    def test_the_body_rejects_unknown_fields_and_tenant_ids(self, ctx: Ctx) -> None:
        h = ctx.admin()
        assert (
            ctx.client.post(
                f"{M}/entities/vehicle_type", json={"values": TYPE, "tenant_id": "t-b"}, headers=h
            ).status_code
            == 422
        )
        response = ctx.client.post(
            f"{M}/entities/vehicle_type", json={"values": {**TYPE, "tenant_id": "t-b"}}, headers=h
        )
        assert response.status_code == 422 and "no existe" in _msg(response)

    def test_listing_exposes_rows_and_a_truncation_flag(self, ctx: Ctx) -> None:
        body = ctx.client.get(f"{M}/entities/daily_demand", headers=ctx.admin()).json()
        assert body["total"] > len(body["rows"]) == 500 and body["truncated"] is True


class TestTenantIsolation:
    def test_changes_in_one_company_never_reach_the_other(self, ctx: Ctx) -> None:
        a, b = ctx.admin(A), ctx.admin(B)
        ctx.client.post(f"{M}/entities/vehicle_type", json={"values": TYPE}, headers=a)
        assert "furgon" in {
            r["name"]
            for r in ctx.client.get(f"{M}/entities/vehicle_type", headers=a).json()["rows"]
        }
        assert "furgon" not in {
            r["name"]
            for r in ctx.client.get(f"{M}/entities/vehicle_type", headers=b).json()["rows"]
        }
        assert ctx.client.delete(f"{M}/entities/vehicle_type/furgon", headers=b).status_code == 404

    def test_history_and_parameters_are_per_company(self, ctx: Ctx) -> None:
        a, b = ctx.admin(A), ctx.admin(B)
        ctx.client.put(f"{M}/parameters/reserve", json={"value": 20}, headers=a)

        def value(h: dict[str, str]) -> object:
            rows = ctx.client.get(f"{M}/parameters", headers=h).json()
            return next(p for p in rows if p["key"] == "reserve")["value"]

        assert value(a) == 20 and value(b) == 8
        assert ctx.client.get(f"{M}/history", headers=b).json() == []

    def test_each_company_has_its_own_seeded_data(self, ctx: Ctx) -> None:
        a = ctx.client.get(f"{M}/entities/customer", headers=ctx.admin(A)).json()["total"]
        b = ctx.client.get(f"{M}/entities/customer", headers=ctx.admin(B)).json()["total"]
        assert (a, b) == (8, 5)


class TestCsvOverHttp:
    def upload(self, ctx: Ctx, kind: str, text: str, tenant: str = A):  # type: ignore[no-untyped-def]
        return ctx.client.post(
            f"{M}/entities/customer/csv/{kind}", json={"content": text}, headers=ctx.admin(tenant)
        )

    def test_preview_then_apply(self, ctx: Ctx) -> None:
        text = CUSTOMER_HEADER + "\nC-901,Uno,Calle 1,4.65,-74.1,15,08:00,12:00,\n"
        preview = self.upload(ctx, "preview", text).json()
        assert preview["can_apply"] and preview["to_create"] == 1
        applied = self.upload(ctx, "apply", text)
        assert applied.status_code == 200 and applied.json() == {
            "entity": "customer",
            "created": 1,
            "updated": 0,
        }
        rows = ctx.client.get(f"{M}/entities/customer", headers=ctx.admin()).json()["rows"]
        assert "C-901" in {r["code"] for r in rows}
        assert any(e.action == "master.csv.apply" for e in ctx.app.audit.list_events(None, 50))

    def test_errors_block_the_load_and_say_where(self, ctx: Ctx) -> None:
        text = CUSTOMER_HEADER + "\nC-901,Uno,Calle 1,40,-74.1,15,08:00,12:00,\n"
        preview = self.upload(ctx, "preview", text).json()
        assert not preview["can_apply"] and preview["errors"][0]["row"] == 2
        response = self.upload(ctx, "apply", text)
        assert response.status_code == 422 and "no se aplicó ninguna fila" in _msg(response)

    def test_a_large_csv_is_accepted_up_to_its_own_limit_and_rejected_beyond(
        self, ctx: Ctx
    ) -> None:
        h = ctx.admin()
        rows = "\n".join(
            f"URB-N,{2000 + i // 365}-{(i % 365) // 31 + 1:02d}-{(i % 28) + 1:02d},5"
            for i in range(4000)
        )
        body = {"content": "corridor_code,date,tons\n" + rows}
        assert len(str(body)) > 64 * 1024  # más que el límite general de cuerpo
        assert (
            ctx.client.post(
                f"{M}/entities/daily_demand/csv/preview", json=body, headers=h
            ).status_code
            == 200
        )
        huge = {"content": "x" * 300_000}
        assert ctx.client.post(
            f"{M}/entities/customer/csv/preview", json=huge, headers=h
        ).status_code in (413, 422)

    def test_other_json_endpoints_keep_the_small_body_limit(self, ctx: Ctx) -> None:
        big = {"values": {"name": "x" * 70_000}}
        response = ctx.client.post(f"{M}/entities/vehicle_type", json=big, headers=ctx.admin())
        assert response.status_code == 413

    def test_the_company_of_the_user_is_the_only_one_affected(self, ctx: Ctx) -> None:
        text = CUSTOMER_HEADER + "\nC-901,Uno,Calle 1,4.65,-74.1,15,08:00,12:00,\n"
        self.upload(ctx, "apply", text, A)
        rows_b = ctx.client.get(f"{M}/entities/customer", headers=ctx.admin(B)).json()["rows"]
        assert "C-901" not in {r["code"] for r in rows_b}

    def test_csv_upload_is_rate_limited(self, ctx: Ctx) -> None:
        text = CUSTOMER_HEADER + "\nC-901,Uno,Calle 1,4.65,-74.1,15,08:00,12:00,\n"
        h = ctx.admin()  # un solo inicio de sesión: el límite medido es el de las cargas
        statuses = [
            ctx.client.post(
                f"{M}/entities/customer/csv/preview", json={"content": text}, headers=h
            ).status_code
            for _ in range(32)
        ]
        assert 429 in statuses and statuses.count(200) == 30


class TestParametersOverHttp:
    def test_set_with_effective_date_and_read_the_versions(self, ctx: Ctx) -> None:
        h = ctx.admin()
        response = ctx.client.put(
            f"{M}/parameters/design_percentile",
            json={"value": 85, "note": "ajuste anual"},
            headers=h,
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["value"] == 85.0 and len(body["versions"]) == 2
        assert body["versions"][0]["in_force"] and body["versions"][1]["value"] == 90.0

    @pytest.mark.parametrize(
        ("key", "payload", "status"),
        [
            ("reserve", {"value": 120}, 422),
            ("reserve", {"value": 9, "valid_from": "2020-01-01"}, 422),
            ("reserve", {"value": 9, "unknown": 1}, 422),
            ("inventado", {"value": 9}, 404),
            ("shift_end", {"value": "05:00"}, 422),
        ],
    )
    def test_invalid_parameter_changes_are_rejected(
        self, ctx: Ctx, key: str, payload: dict, status: int
    ) -> None:  # type: ignore[type-arg]
        assert (
            ctx.client.put(f"{M}/parameters/{key}", json=payload, headers=ctx.admin()).status_code
            == status
        )

    def test_a_parameter_change_moves_the_capacity_dashboard(self, ctx: Ctx) -> None:
        manager = ctx.headers(A, "gerencia")
        before = ctx.client.get("/api/v1/capacity/dashboard", headers=manager).json()["totals"][
            "balance"
        ]
        ctx.client.put(f"{M}/parameters/reserve", json={"value": 40}, headers=ctx.admin())
        after = ctx.client.get("/api/v1/capacity/dashboard", headers=manager).json()["totals"][
            "balance"
        ]
        assert after > before

    def test_a_fleet_change_moves_the_gap(self, ctx: Ctx) -> None:
        manager = ctx.headers(A, "gerencia")

        def gap() -> int:
            body = ctx.client.get("/api/v1/capacity/dashboard", headers=manager).json()
            return body["totals"]["gap"]  # type: ignore[no-any-return]

        before = gap()
        ctx.client.put(
            f"{M}/entities/fleet/turbo",
            json={"values": {"vehicle_type": "turbo", "count": 60}},
            headers=ctx.admin(),
        )
        assert gap() > before


class TestExports:
    def test_master_export_as_csv_and_xlsx(self, ctx: Ctx) -> None:
        h = ctx.admin()
        csv = ctx.client.get(f"{M}/entities/vehicle_type/export?format=csv", headers=h)
        assert csv.status_code == 200 and csv.headers["content-type"].startswith("text/csv")
        assert 'filename="maestro-vehicle_type.csv"' in csv.headers["content-disposition"]
        assert csv.content.startswith(b"\xef\xbb\xbf") and b"turbo" in csv.content
        xlsx = ctx.client.get(f"{M}/entities/vehicle_type/export?format=xlsx", headers=h)
        ws = load_workbook(io.BytesIO(xlsx.content)).active
        assert ws["A1"].value == "DATOS DE DEMOSTRACIÓN"
        assert any(ws.cell(row=r, column=1).value == "turbo" for r in range(1, 12))

    def test_an_unsupported_format_is_rejected_and_exports_are_audited(self, ctx: Ctx) -> None:
        h = ctx.admin()
        assert (
            ctx.client.get(f"{M}/entities/customer/export?format=exe", headers=h).status_code == 422
        )
        ctx.client.get(f"{M}/entities/customer/export?format=csv", headers=h)
        assert any(e.action == "master.export" for e in ctx.app.audit.list_events(None, 50))

    def test_capacity_dashboard_export_to_excel(self, ctx: Ctx) -> None:
        response = ctx.client.get(
            "/api/v1/capacity/export.xlsx", headers=ctx.headers(A, "gerencia")
        )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        wb = load_workbook(io.BytesIO(response.content))
        assert wb.sheetnames == ["Flota por método", "Corredores", "Proyección", "Indicadores"]
        assert wb["Flota por método"]["A1"].value == "DATOS DE DEMOSTRACIÓN"
        assert any(e.action == "export.generate" for e in ctx.app.audit.list_events(None, 50))

    def test_capacity_export_follows_the_capacity_permission(self, ctx: Ctx) -> None:
        for user in ("despachador", "conductor"):
            assert (
                ctx.client.get(
                    "/api/v1/capacity/export.xlsx", headers=ctx.headers(A, user)
                ).status_code
                == 403
            )
        assert ctx.client.get("/api/v1/capacity/export.xlsx").status_code == 401
        assert (
            ctx.client.get("/api/v1/capacity/export.xlsx", headers=ctx.admin()).status_code == 200
        )

    def test_driver_route_export_contains_only_their_vehicle(self, ctx: Ctx) -> None:
        h = ctx.headers(A, "conductor")
        csv = ctx.client.get("/api/v1/routing/driver/route/export?format=csv", headers=h)
        assert csv.status_code == 200 and b"T-01" in csv.content
        xlsx = ctx.client.get("/api/v1/routing/driver/route/export?format=xlsx", headers=h)
        ws = load_workbook(io.BytesIO(xlsx.content)).active
        assert ws["A2"].value == "Vehículo: T-01"
        assert "attachment" in xlsx.headers["content-disposition"]
        assert (
            ctx.client.get(
                "/api/v1/routing/driver/route/export", headers=ctx.headers(A, "despachador")
            ).status_code
            == 403
        )

    def test_filenames_are_safe(self) -> None:
        from fleetplanner.adapters.inbound.http.routes_masters import safe_filename

        assert safe_filename("../../etc/passwd", "csv") == "etc-passwd.csv"
        assert safe_filename('a"b\r\nc', "xlsx") == "a-b-c.xlsx"
        assert safe_filename("", "csv") == "exportacion.csv"
